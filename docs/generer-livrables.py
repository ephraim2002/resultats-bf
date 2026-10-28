#!/usr/bin/env python3
"""Construit les documents Word du projet à partir des sources Markdown internes.

Usage : python docs/generer-livrables.py [DOSSIER_DE_SORTIE]
Dépendance : python-docx. Les images des logos officiels doivent être fournies
localement sous docs/logo-ibam.png et docs/logo-ujkz.png avant la génération.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "livrables"

NAVY = "16324F"
BLUE = "175CD3"
TEAL = "079E89"
PALE_BLUE = "EAF1FB"
LIGHT = "F2F5F9"
DARK = "263448"
GOLD = "D69E22"
WHITE = "FFFFFF"

TEAM = [
    ("BARRY", "El Nour", "66806583", "SI"),
    ("BAMBARA", "Mireille Epiphanie Joëlle Guerla", "6539632", "SIE"),
    ("SAWADOGO", "Wenden Kogondé Aimé Ephraïm", "72620579", "SI"),
    ("WEDRAOGO", "Relwende Diane Micheline", "76677438", "SIE"),
]


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell_text(cell, text: str, *, bold=False, color=DARK, size=9, center=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    set_cell_margins(cell)


def set_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    props.append(color)
    under = OxmlElement("w:u")
    under.set(qn("w:val"), "single")
    props.append(under)
    run.append(props)
    txt = OxmlElement("w:t")
    txt.text = text
    run.append(txt)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instr, separate, value, end):
        run._r.append(item)


def add_running_matter(doc: Document, short_title: str):
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.paragraph_format.space_after = Pt(2)
        r = header.add_run("GROUPE 18  ·  RÉSULTATS-BF")
        r.font.name = "Aptos"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(BLUE)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(2)
        rf = footer.add_run(short_title + "  ·  page ")
        rf.font.name = "Aptos"
        rf.font.size = Pt(8)
        rf.font.color.rgb = RGBColor.from_string("617186")
        set_page_field(footer)


def add_logo_row(doc: Document):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(2.0)
    for cell, label, path in (
        (table.cell(0, 0), "IBAM", DOCS / "logo-ibam.png"),
        (table.cell(0, 1), "UJKZ", DOCS / "logo-ujkz.png"),
    ):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        if path.is_file():
            p.add_run().add_picture(str(path), width=Inches(1.18), height=Inches(1.18))
        else:
            run = p.add_run(label)
            run.bold = True
            run.font.size = Pt(17)
            run.font.color.rgb = RGBColor.from_string(NAVY)
            p2 = cell.add_paragraph("Insérer le logo officiel")
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(0)
            for r in p2.runs:
                r.italic = True
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor.from_string("66758A")


def add_team_table(doc: Document, compact=False):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ("Nom", "Prénom", "Numéro", "Filière")
    for i, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=WHITE, size=8, center=True)
        shade(table.rows[0].cells[i], NAVY)
    set_repeat_table_header(table.rows[0])
    for idx, member in enumerate(TEAM):
        row = table.add_row()
        for col, text in enumerate(member):
            set_cell_text(row.cells[col], text, size=7.5 if compact else 8, center=(col >= 2))
            shade(row.cells[col], WHITE if idx % 2 else LIGHT)
    widths = (1.12, 3.60, 1.05, 0.70) if not compact else (1.02, 3.30, 0.93, 0.65)
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
    return table


def add_cover(doc: Document, *, title: str, subtitle: str, short_label: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("UNIVERSITÉ JOSEPH KI-ZERBO")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("INSTITUT BURKINABÈ DES ARTS ET MÉTIERS · IBAM")
    r.font.size = Pt(10)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    add_logo_row(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title.upper())
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(24 if short_label == "Dossier d’architecture" else 22)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(subtitle)
    r.font.name = "Aptos"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string("53657A")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("PROJET DE GROUPE 18  ·  4 MEMBRES")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Master 1 Informatique  ·  Virtualisation et Cloud Computing")
    r.font.size = Pt(9)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("Année universitaire 2026–2027  ·  Enseignant : Dr SOMDA Flavien Hervé")
    r.font.size = Pt(9)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ÉQUIPE DU PROJET")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    add_team_table(doc, compact=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    p.add_run("Dépôt Git : ")
    set_hyperlink(p, "github.com/ephraim2002/resultats-bf", "https://github.com/ephraim2002/resultats-bf")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Document de travail daté du 27 octobre 2026 · Statut des scénarios détaillé dans le rapport")
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string("6C7A8B")


def set_cell_borders(cell, color: str = "C9D4E2"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def add_architecture_figure(doc: Document):
    image_path = DOCS / "cahier-figure.png"
    if not image_path.is_file():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(image_path), width=Inches(6.22))
    caption = doc.add_paragraph("Figure 1 — Architecture cible fournie dans le cahier des charges du projet.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(4)
    for r in caption.runs:
        r.italic = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string("5D6C7F")


def add_rich_text(paragraph, text: str, font_size=9):
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.font.size = Pt(font_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.size = Pt(font_size)
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Cascadia Code"
            run.font.size = Pt(max(font_size - 1, 7))
            run.font.color.rgb = RGBColor.from_string(BLUE)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.font.size = Pt(font_size)


def add_table_from_lines(doc: Document, lines: list[str]):
    parsed = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in parsed if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in row)]
    if not rows:
        return
    cols = max(map(len, rows))
    table = doc.add_table(rows=1, cols=cols)
    table.autofit = True
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for idx in range(cols):
        text = rows[0][idx] if idx < len(rows[0]) else ""
        set_cell_text(table.rows[0].cells[idx], text, bold=True, color=WHITE, size=7.7)
        shade(table.rows[0].cells[idx], NAVY)
        set_cell_borders(table.rows[0].cells[idx])
    for ri, data in enumerate(rows[1:]):
        row = table.add_row()
        for ci in range(cols):
            value = data[ci] if ci < len(data) else ""
            set_cell_text(row.cells[ci], value, size=7.1)
            shade(row.cells[ci], WHITE if ri % 2 else LIGHT)
            set_cell_borders(row.cells[ci])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, lines: Iterable[str]):
    lines = list(lines)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_together = True
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT)
    pPr.append(shd)
    for i, line in enumerate(lines):
        run = p.add_run(line if line else " ")
        run.font.name = "Cascadia Code"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(DARK)
        if i < len(lines) - 1:
            run.add_break()


def add_source_document(doc: Document, source: Path, *, page_break_tokens=False):
    text = source.read_text(encoding="utf-8-sig")
    content_pages = text.split("---PAGE---") if page_break_tokens else [text]
    for page_idx, page in enumerate(content_pages):
        if page_idx <= 1:
            # The split starts with an empty segment and then the source cover;
            # the formatted cover has already been built by add_cover().
            continue
        doc.add_page_break()
        lines = page.splitlines()
        if page_idx == 0 and lines and lines[0].startswith("# "):
            lines = lines[1:]
        if page_idx == 1 and lines:
            pass
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            if not line.strip():
                i += 1
                continue
            if line.startswith("[Emplacement réservé au logo officiel") or line.startswith("NOTE SUR LES ÉLÉMENTS GRAPHIQUES"):
                i += 1
                continue
            if line.strip() == "TABLE DES MATIÈRES":
                p = doc.add_paragraph(style="Heading 1")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                add_rich_text(p, line.strip(), 15)
            elif line.strip().startswith("[Emplacement pour le logo officiel"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(2)
                p.add_run(line.strip().strip("[]"))
            elif (
                len(line) < 100
                and line == line.upper()
                and re.search(r"[A-ZÉÈÀÂÊÎÔÛÇ]", line)
                and not line.startswith("[Emplacement")
            ):
                p = doc.add_paragraph(style="Heading 1")
                p.paragraph_format.keep_with_next = True
                add_rich_text(p, line, 13)
                if "VUE D’ENSEMBLE" in line.upper() or "VUE D'ENSEMBLE" in line.upper():
                    add_architecture_figure(doc)
            elif line.startswith("# "):
                p = doc.add_paragraph(style="Title")
                p.paragraph_format.keep_with_next = True
                add_rich_text(p, line[2:], 22)
            elif line.startswith("## "):
                p = doc.add_paragraph(style="Heading 1")
                p.paragraph_format.keep_with_next = True
                add_rich_text(p, line[3:], 15)
                if "VUE D’ENSEMBLE" in line.upper() or "VUE D'ENSEMBLE" in line.upper():
                    add_architecture_figure(doc)
            elif line.startswith("### "):
                p = doc.add_paragraph(style="Heading 2")
                p.paragraph_format.keep_with_next = True
                add_rich_text(p, line[4:], 12)
            elif line.strip().startswith("|"):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    table_lines.append(lines[i])
                    i += 1
                add_table_from_lines(doc, table_lines)
                continue
            elif line.strip().startswith("•") or line.strip().startswith("- "):
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                add_rich_text(p, line.strip()[2:] if line.strip().startswith("- ") else line.strip()[1:].strip(), 8.5)
            elif re.match(r"^\d+[.)]\s", line):
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(2)
                add_rich_text(p, re.sub(r"^\d+[.)]\s", "", line), 8.5)
            elif line.startswith("    "):
                code_lines = []
                while i < len(lines) and (lines[i].startswith("    ") or not lines[i].strip()):
                    code_lines.append(lines[i][4:] if lines[i].startswith("    ") else "")
                    i += 1
                add_code_block(doc, code_lines)
                continue
            elif line.startswith("|"):
                add_table_from_lines(doc, [line])
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1.5)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.widow_control = True
                add_rich_text(p, line, 8.2)
            i += 1


def setup_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(2)
    for name, size, color in (("Title", 24, NAVY), ("Heading 1", 15, NAVY), ("Heading 2", 11, BLUE), ("Heading 3", 10, TEAL)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(4)


def make_doc(source: Path, outfile: Path, title: str, subtitle: str, *, page_tokens=False):
    doc = Document()
    setup_styles(doc)
    add_running_matter(doc, title)
    add_cover(doc, title=title, subtitle=subtitle, short_label=title)
    add_source_document(doc, source, page_break_tokens=page_tokens)
    doc.core_properties.title = title + " — RÉSULTATS-BF — Groupe 18"
    doc.core_properties.subject = subtitle
    doc.core_properties.author = "Groupe 18 — UJKZ / IBAM"
    doc.core_properties.keywords = "RÉSULTATS-BF; Vagrant; VirtualBox; infrastructure; groupe 18"
    outfile.parent.mkdir(parents=True, exist_ok=True)
    doc.save(outfile)


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    make_doc(
        DOCS / "architecture.md",
        OUTPUT / "Dossier-architecture-RESULTATS-BF.docx",
        "Dossier d’architecture",
        "Capacité cible : quatre serveurs applicatifs · Exécution mesurée : deux, selon les ressources disponibles",
        page_tokens=True,
    )
    make_doc(
        DOCS / "journal-de-bord.md",
        OUTPUT / "Journal-de-bord-RESULTATS-BF.docx",
        "Journal de bord",
        "Travaux réalisés, contrôles mesurés et limites des essais",
        page_tokens=True,
    )
    print("Documents générés dans :", OUTPUT)
    for path in sorted(OUTPUT.glob("*.docx")):
        print(f"  {path.name} — {path.stat().st_size:,} octets")
    if not (DOCS / "logo-ibam.png").is_file() or not (DOCS / "logo-ujkz.png").is_file():
        print("Logos officiels non intégrés : aucun fichier officiel autorisé n’a été fourni ; les documents affichent des emplacements explicites.")


if __name__ == "__main__":
    main()
